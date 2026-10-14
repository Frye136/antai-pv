"""Référentiel clients / véhicules / contrats : définition des champs, import Excel/CSV, modèles, exports."""
import csv
import re
import unicodedata
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from . import db, parser, plates

# (colonne, libellé, type)  — types : text, date, datetime, bool, number, type_client, vehicule, client
FIELDS = {
    "clients": [
        ("code_client", "Code client", "text"),
        ("type", "Type (PM/PP)", "type_client"),
        ("raison_sociale", "Raison sociale", "text"),
        ("siret", "SIRET", "text"),
        ("civilite", "Civilité", "text"),
        ("nom", "Nom", "text"),
        ("prenom", "Prénom", "text"),
        ("date_naissance", "Date de naissance", "date"),
        ("lieu_naissance", "Lieu de naissance", "text"),
        ("num_permis", "N° permis", "text"),
        ("date_permis", "Permis délivré le", "date"),
        ("adresse", "Adresse", "text"),
        ("complement", "Complément", "text"),
        ("code_postal", "Code postal", "text"),
        ("ville", "Ville", "text"),
        ("pays", "Pays", "text"),
        ("email", "E-mail", "text"),
        ("email_cc", "E-mail copie", "text"),
        ("telephone", "Téléphone", "text"),
        ("frais_dossier", "Frais de dossier spécifiques", "number"),
        ("traitement_pv", "Contraventions : payer et refacturer", "traitement"),
        ("notes", "Notes", "text"),
        ("actif", "Actif", "bool"),
    ],
    "vehicules": [
        ("immatriculation", "Immatriculation", "text"),
        ("marque", "Marque", "text"),
        ("modele", "Modèle", "text"),
        ("categorie", "Catégorie", "text"),
        ("vin", "N° de série (VIN)", "text"),
        ("notes", "Notes", "text"),
        ("actif", "Actif", "bool"),
    ],
    "contrats": [
        ("numero", "N° contrat", "text"),
        ("vehicule_id", "Immatriculation", "vehicule"),
        ("client_id", "Code client", "client"),
        ("date_debut", "Début (date heure)", "datetime"),
        ("date_fin", "Fin (date heure)", "datetime"),
        ("cond_nom", "Conducteur nom", "text"),
        ("cond_prenom", "Conducteur prénom", "text"),
        ("cond_date_naissance", "Conducteur date naissance", "date"),
        ("cond_lieu_naissance", "Conducteur lieu naissance", "text"),
        ("cond_num_permis", "Conducteur n° permis", "text"),
        ("cond_adresse", "Conducteur adresse", "text"),
        ("cond_code_postal", "Conducteur CP", "text"),
        ("cond_ville", "Conducteur ville", "text"),
        ("designer_conducteur", "Désigner le conducteur", "bool"),
        ("notes", "Notes", "text"),
    ],
}
KEYS = {"clients": "code_client", "vehicules": "immatriculation", "contrats": "numero"}

# libellés alternatifs reconnus à l'import (export ERP / logiciel de location)
ALIASES = {
    "code_client": ["code", "n client", "num client", "numero client", "ref client", "id client", "client"],
    "raison_sociale": ["societe", "entreprise", "denomination", "nom societe"],
    "siret": ["siren", "siret siren"],
    "nom": ["nom client", "nom de famille"],
    "date_naissance": ["ne le", "naissance"],
    "num_permis": ["permis", "numero permis", "n permis de conduire"],
    "adresse": ["adresse 1", "rue", "adresse ligne 1", "voie"],
    "complement": ["adresse 2", "complement adresse"],
    "code_postal": ["cp", "code postal"],
    "ville": ["commune", "localite"],
    "email": ["mail", "courriel", "adresse mail", "e mail"],
    "immatriculation": ["immat", "plaque", "immatriculation vehicule", "vehicule"],
    "numero": ["contrat", "n contrat", "numero contrat", "num contrat", "ref contrat"],
    "date_debut": ["debut", "date depart", "depart", "date debut location", "date de debut"],
    "date_fin": ["fin", "date retour", "retour", "date fin location", "date de fin"],
}


def _key(s: str) -> str:
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def _map_headers(table: str, headers: list[str]) -> dict[int, str]:
    lookup = {}
    for col, label, _ in FIELDS[table]:
        for name in [col, label] + ALIASES.get(col, []):
            lookup.setdefault(_key(name), col)
    if table == "contrats":  # dans un fichier de contrats, « client » = code client
        lookup[_key("client")] = "client_id"
        lookup[_key("code client")] = "client_id"
        lookup[_key("immatriculation")] = "vehicule_id"
        for a in ALIASES["immatriculation"]:
            lookup[_key(a)] = "vehicule_id"
    return {i: lookup[_key(h)] for i, h in enumerate(headers) if _key(h) in lookup}


def read_table(path: str) -> tuple[list[str], list[list]]:
    p = Path(path)
    if p.suffix.lower() in (".xlsx", ".xlsm"):
        wb = load_workbook(p, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
    else:
        raw = p.read_bytes()
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("cp1252")
        first = text.splitlines()[0] if text else ""
        sep = max([";", ",", "\t"], key=first.count)
        rows = list(csv.reader(text.splitlines(), delimiter=sep))
    # première ligne non vide = en-têtes
    while rows and not any(c not in (None, "") for c in rows[0]):
        rows.pop(0)
    if not rows:
        return [], []
    return [str(h or "").strip() for h in rows[0]], rows[1:]


def norm_datetime(v, end: bool = False) -> str | None:
    """Valeur Excel/texte -> 'YYYY-MM-DD HH:MM' (fin de journée si `end` et pas d'heure)."""
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        if end and v.hour == 0 and v.minute == 0:
            return v.strftime("%Y-%m-%d 23:59")
        return v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, date):
        return v.isoformat() + (" 23:59" if end else " 00:00")
    s = str(v).strip()
    for f in ("%d/%m/%Y %H:%M", "%d/%m/%Y %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S", "%d/%m/%Y %Hh%M",
              "%d-%m-%Y %H:%M"):
        try:
            return datetime.strptime(s, f).strftime("%Y-%m-%d %H:%M")
        except ValueError:
            pass
    d = parser.to_date(s)
    return (d.isoformat() + (" 23:59" if end else " 00:00")) if d else None


def norm_date(v) -> str | None:
    d = parser.to_date(v.date() if isinstance(v, datetime) else v)
    return d.isoformat() if d else None


def coerce(table: str, col: str, kind: str, v):
    """Conversion d'une valeur saisie / importée vers le format stocké."""
    if isinstance(v, str):
        v = v.strip()
    if v in (None, ""):
        return None
    if kind == "date":
        return norm_date(v)
    if kind == "datetime":
        return norm_datetime(v, end=(col == "date_fin"))
    if kind == "bool":
        return 1 if _key(v) in ("1", "oui", "o", "x", "vrai", "true", "yes") or v is True or v == 1 else 0
    if kind == "number":
        try:
            return float(str(v).replace(",", ".").replace("€", "").strip())
        except ValueError:
            return None
    if kind == "traitement":  # « oui » / PAIEMENT = on paie et on refacture au lieu de désigner
        return "PAIEMENT" if _key(v) in ("1", "oui", "o", "x", "paiement", "payer", "refacturer", "true") else None
    if kind == "type_client":
        k = _key(v)
        return "PP" if k in ("pp", "particulier", "personne physique", "physique") else "PM"
    if col == "immatriculation":
        return plates.normalize(v)
    if col == "code_postal":
        s = str(v).split(".")[0]
        return s.zfill(5) if s.isdigit() and len(s) == 4 else s
    if col in ("code_client", "numero", "siret", "num_permis") and isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def import_file(table: str, path: str) -> dict:
    """Import / mise à jour (clé : code client, immatriculation, n° contrat). Renvoie un bilan."""
    headers, rows = read_table(path)
    mapping = _map_headers(table, headers)
    key = KEYS[table]
    if table == "contrats":
        if not {"vehicule_id", "client_id", "date_debut"} <= set(mapping.values()):
            raise ValueError("Colonnes obligatoires : Immatriculation, Code client, Début.")
    elif key not in mapping.values():
        raise ValueError(f"Colonne « {dict((c, l) for c, l, _ in FIELDS[table])[key]} » introuvable.")
    kinds = {c: k for c, _, k in FIELDS[table]}
    created = updated = 0
    errors = []
    conn = db.get()
    with db.lock:
        for n, row in enumerate(rows, 2):
            if not any(c not in (None, "") for c in row):
                continue
            vals = {}
            for i, col in mapping.items():
                if i < len(row):
                    vals[col] = row[i]
            try:
                if table == "contrats":
                    immat = plates.normalize(str(vals.pop("vehicule_id") or ""))
                    veh = db.q1("SELECT id FROM vehicules WHERE immatriculation=?", (immat,))
                    if not veh:  # véhicule inconnu : créé à la volée
                        conn.execute("INSERT INTO vehicules(immatriculation) VALUES(?)", (immat,))
                        veh = db.q1("SELECT id FROM vehicules WHERE immatriculation=?", (immat,))
                    code = coerce("clients", "code_client", "text", vals.pop("client_id"))
                    cli = db.q1("SELECT id FROM clients WHERE code_client=?", (code,))
                    if not cli:
                        raise ValueError(f"client « {code} » inconnu (importez les clients d'abord)")
                    vals = {c: coerce(table, c, kinds[c], v) for c, v in vals.items()}
                    vals["vehicule_id"], vals["client_id"] = veh["id"], cli["id"]
                    if not vals.get("date_debut"):
                        raise ValueError("date de début invalide")
                    if not vals.get("numero"):
                        vals["numero"] = f"{immat}-{vals['date_debut'][:10]}"
                else:
                    vals = {c: coerce(table, c, kinds[c], v) for c, v in vals.items()}
                    if not vals.get(key):
                        raise ValueError(f"{key} vide")
                existing = db.q1(f"SELECT id FROM {table} WHERE {key}=?", (vals[key],))
                if existing:
                    cols = ", ".join(f"{c}=?" for c in vals)
                    conn.execute(f"UPDATE {table} SET {cols} WHERE id=?", (*vals.values(), existing["id"]))
                    updated += 1
                else:
                    cols = ", ".join(vals)
                    conn.execute(f"INSERT INTO {table}({cols}) VALUES({', '.join('?' * len(vals))})",
                                 tuple(vals.values()))
                    created += 1
            except Exception as e:
                errors.append(f"ligne {n} : {e}")
    conn.commit()
    db.log(f"Import {table}", f"{Path(path).name} : {created} créés, {updated} mis à jour, {len(errors)} erreurs")
    return {"created": created, "updated": updated, "errors": errors,
            "colonnes": [headers[i] + " → " + c for i, c in mapping.items()]}


def _style_header(ws, ncols):
    for i in range(1, ncols + 1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3A5F")
        ws.column_dimensions[get_column_letter(i)].width = 20
    ws.freeze_panes = "A2"


def write_model(table: str, path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = table
    labels = [l for _, l, _ in FIELDS[table]]
    ws.append(labels)
    examples = {
        "clients": ["C001", "PM", "TRANSPORTS DUPONT", "12345678900012", "", "", "", "", "", "", "",
                    "12 rue de la Gare", "", "55100", "VERDUN", "FRANCE", "compta@dupont.example.com", "", "",
                    "", "", "1"],
        "vehicules": ["AB-123-CD", "RENAULT", "MASTER", "Utilitaire", "", "", "1"],
        "contrats": ["LOC-2026-001", "AB-123-CD", "C001", "01/10/2026 08:00", "15/10/2026 18:00",
                     "", "", "", "", "", "", "", "", "0", ""],
    }
    ws.append(examples[table])
    _style_header(ws, len(labels))
    wb.save(path)


def export_pv_excel(rows, path: str) -> None:
    """Liste des PV (tels qu'affichés) vers Excel."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Contraventions"
    cols = ["N° avis", "Type", "Date avis", "Date infraction", "Heure", "Immatriculation", "Lieu", "Nature",
            "Montant", "Client", "Contrat", "Filière", "Statut", "Échéance", "Dénoncé le", "Payé le", "Montant payé",
            "Mode de paiement", "Réf. paiement", "N° facture", "Facturé le", "Réglé le", "Notifié le",
            "Frais dossier", "Frais refacturés", "Anomalies"]
    ws.append(cols)
    for r in rows:
        ws.append([r["numero_avis"], r["type"], parser.to_date(r["date_avis"]), parser.to_date(r["date_infraction"]),
                   r["heure_infraction"], plates.display(r["immatriculation"]), r["lieu"], r["nature"], r["montant"],
                   r["client"], r["contrat"], db.TRAITEMENTS.get(r["traitement"], ""),
                   db.STATUTS.get(r["statut"], r["statut"]),
                   parser.to_date(r["date_limite"]), parser.to_date(r["date_designation"]),
                   parser.to_date(r["date_paiement"]), r["montant_paye"], r["mode_paiement"], r["ref_paiement"],
                   r["facture_numero"], parser.to_date(r["date_facture"]), parser.to_date(r["date_reglement"]),
                   r["date_notification"], r["frais_dossier"], "Oui" if r["refacture"] else "Non", r["anomalie"]])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, date):
                c.number_format = "DD/MM/YYYY"
    _style_header(ws, len(cols))
    wb.save(path)


def export_refacturation(path: str, mark: bool) -> int:
    """À facturer, par client : frais de dossier des avis dénoncés + amendes payées pour le compte du client
    (débours) et leurs frais. `mark` : frais marqués facturés, avis payés -> « Refacturé – attente règlement »."""
    rows = db.q("""SELECT pv.*, c.code_client, COALESCE(c.raison_sociale, c.nom||' '||COALESCE(c.prenom,'')) client,
                          k.numero contrat
                   FROM pv JOIN clients c ON c.id=pv.client_id LEFT JOIN contrats k ON k.id=pv.contrat_id
                   WHERE (pv.statut='PAYE')
                      OR (pv.refacture=0 AND pv.statut IN ('EXPORTE','DEPOSE','CLOS') AND pv.traitement IS NOT 'PAIEMENT')
                   ORDER BY client, pv.date_infraction""")
    wb = Workbook()
    ws = wb.active
    ws.title = "Détail"
    cols = ["Code client", "Client", "Contrat", "N° avis", "Type", "Date infraction", "Immatriculation", "Filière",
            "Amende réglée pour le client (débours)", "Payée le", "Frais de dossier HT", "Total à facturer"]
    ws.append(cols)
    totals: dict = {}
    for r in rows:
        frais = r["frais_dossier"] or 0
        debours = (r["montant_paye"] or 0) if r["statut"] == "PAYE" else 0
        total = frais + debours
        ws.append([r["code_client"], r["client"], r["contrat"], r["numero_avis"], r["type"],
                   parser.to_date(r["date_infraction"]), plates.display(r["immatriculation"]),
                   db.TRAITEMENTS.get(r["traitement"], ""), debours or None, parser.to_date(r["date_paiement"]),
                   frais, total])
        t = totals.setdefault((r["code_client"], r["client"]), [0, 0.0, 0.0])
        t[0] += 1
        t[1] += debours
        t[2] += frais
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, date):
                c.number_format = "DD/MM/YYYY"
    _style_header(ws, len(cols))
    ws2 = wb.create_sheet("Synthèse")
    ws2.append(["Code client", "Client", "Nb avis", "Débours (amendes réglées)", "Frais de dossier HT", "Total"])
    for (code, name), (n, deb, fr) in sorted(totals.items(), key=lambda x: str(x[0][1])):
        ws2.append([code, name, n, round(deb, 2), round(fr, 2), round(deb + fr, 2)])
    _style_header(ws2, 6)
    wb.save(path)
    if mark:
        for r in rows:
            upd = {"refacture": 1}
            if r["statut"] == "PAYE":
                upd["statut"] = "REFACTURE"
                upd["date_facture"] = date.today().isoformat()
            db.update("pv", r["id"], upd)
        db.log("Refacturation", f"{len(rows)} avis exportés pour facturation ({Path(path).name})")
    return len(rows)
